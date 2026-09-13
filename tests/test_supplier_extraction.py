"""Document reading and the trust guards applied to what the model reports."""
from __future__ import annotations

import os
import tempfile
import unittest

from rfq_copilot import supplier_guards as sg
from rfq_copilot.document_extractor import (
    _pdf_unescape,
    DocumentExtractorRegistry, DocxExtractor, PdfExtractor, TextExtractor, XlsxExtractor,
)
from rfq_copilot.supplier_models import (
    Certification, ClaimStatus, Evidence, ExtractionStatus, MatchStatus, NormalizationStatus,
    PriceBasis, QuestionnaireResponse, QuoteStatus, SourceDocument, SupplierQuote,
)
from tests.supplier_helpers import FIXTURES, carton_rfq, make_documents


def fixture(name):
    return os.path.join(FIXTURES, name)


class DocumentExtractionTest(unittest.TestCase):
    """Each format is read for real; nothing is fabricated for a file we cannot open."""

    def test_xlsx_gives_cell_level_locations(self):
        c = XlsxExtractor().extract(fixture("supplier_a_quote.xlsx"))
        self.assertEqual(c.status, ExtractionStatus.EXTRACTED)
        self.assertIn("0.42", c.text)
        cells = [b for b in c.blocks if b.cell]
        self.assertTrue(cells, "evidence needs addressable cells")
        self.assertTrue(any(b.sheet == "Quotation" for b in cells))
        self.assertTrue(any("cell" in b.location for b in cells))

    def test_a_formula_with_no_cached_result_is_listed_not_dropped(self):
        """A workbook written by a program and never opened in Excel stores no result for
        its formulas. Dropping those cells silently loses a whole column — a quotation
        whose only prices are a computed Total would read as having no prices at all."""
        import openpyxl
        tmp = tempfile.mkdtemp(prefix="xlsx_formula_")
        path = os.path.join(tmp, "quote.xlsx")
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Quotation"
        ws["A1"], ws["B1"], ws["C1"] = "Item", "Qty", "Unit Price"
        ws["D1"] = "Total"
        ws["A2"], ws["B2"], ws["C2"] = "Desk", 120, 18500
        ws["D2"] = "=B2*C2"          # never calculated: openpyxl writes no cached value
        wb.save(path)

        c = XlsxExtractor().extract(path)
        self.assertEqual(c.status, ExtractionStatus.EXTRACTED)
        self.assertIn("18500", c.text, "the values that do exist still come through")
        self.assertIn("=B2*C2", c.text, "the formula cell must not vanish")
        self.assertIn("never stored a result", c.text)
        self.assertIn("formula", c.note)
        self.assertNotIn("2220000", c.text,
                         "the reader must not evaluate a supplier's arithmetic for them")

    def test_pdf_pages_are_separated_and_not_duplicated(self):
        c = PdfExtractor().extract(fixture("supplier_b_quote.pdf"))
        self.assertEqual(c.status, ExtractionStatus.EXTRACTED)
        pages = sorted({b.page for b in c.blocks})
        self.assertEqual(pages, [1, 2, 3], "the 'stream' inside 'endstream' must not open a phantom page")
        page2 = " ".join(b.text for b in c.blocks if b.page == 2)
        page3 = " ".join(b.text for b in c.blocks if b.page == 3)
        self.assertIn("5% discount", page2, "the buried footnote must survive extraction")
        self.assertIn("25 days", page3, "the contradicting lead time is on another page")
        self.assertNotEqual(page2, page3)

    def test_an_ascii85_compressed_pdf_is_read(self):
        """ReportLab — which a great many quotation tools embed — writes
        /Filter [/ASCII85Decode /FlateDecode] by default. Handling only Flate meant those
        PDFs reported "no extractable text operators" and the buyer was told, wrongly,
        that their quotation looked like a scan."""
        from reportlab.pdfgen import canvas
        tmp = tempfile.mkdtemp(prefix="pdf_a85_")
        path = os.path.join(tmp, "quote.pdf")
        c = canvas.Canvas(path)
        c.drawString(72, 800, "Workstation 120 units at INR 19,200 each")
        c.drawString(72, 780, "Lead time 20-25 days. MOQ 25 units.")
        c.save()

        out = PdfExtractor().extract(path)
        self.assertEqual(out.status, ExtractionStatus.EXTRACTED, out.note)
        self.assertIn("19,200", out.text)
        self.assertIn("MOQ 25 units", out.text)

    def test_an_undecodable_pdf_says_why_instead_of_blaming_a_scanner(self):
        """A PDF we cannot decompress is not the same thing as a photograph of a page, and
        telling a supplier the wrong one sends them to fix the wrong problem."""
        body = b"1 0 obj<</Filter/JBIG2Decode/Length 4>>stream\n\x00\x01\x02\x03\nendstream endobj"
        tmp = tempfile.mkdtemp(prefix="pdf_odd_")
        path = os.path.join(tmp, "odd.pdf")
        with open(path, "wb") as f:
            f.write(b"%PDF-1.4\n" + body)
        out = PdfExtractor().extract(path)
        self.assertEqual(out.status, ExtractionStatus.UNSUPPORTED)
        self.assertIn("JBIG2Decode", out.note)

    def test_an_octal_escape_becomes_the_character_it_stands_for(self):
        """A pound or euro sign written as \\243 would otherwise reach the currency guard as
        the literal text "\\243", and be refused as a price it should have accepted."""
        self.assertEqual(_pdf_unescape(rb"(Procurement \226 New Office)"),
                         "Procurement \u2013 New Office")
        self.assertEqual(_pdf_unescape(rb"(\24312.50 per unit)"), "\u00a312.50 per unit")

    def test_docx_numbers_paragraphs(self):
        c = DocxExtractor().extract(fixture("supplier_c_response.docx"))
        self.assertEqual(c.status, ExtractionStatus.EXTRACTED)
        self.assertTrue(any(b.paragraph for b in c.blocks))
        self.assertIn("0.47", c.text)
        self.assertIn("per kg", c.text.lower())

    def test_txt_is_read_directly(self):
        c = TextExtractor().extract(fixture("supplier_d_response.txt"))
        self.assertEqual(c.status, ExtractionStatus.EXTRACTED)
        self.assertIn("41 cents", c.text)

    def test_unknown_format_is_reported_not_invented(self):
        reg = DocumentExtractorRegistry(extractors=[TextExtractor()])
        c = reg.extract(fixture("supplier_a_quote.xlsx"))
        self.assertEqual(c.status, ExtractionStatus.UNSUPPORTED)
        self.assertEqual(c.text, "", "an unreadable document must yield no content at all")

    def test_missing_file_is_reported(self):
        c = DocumentExtractorRegistry().extract(fixture("does_not_exist.pdf"))
        self.assertEqual(c.status, ExtractionStatus.UNSUPPORTED)

    def test_image_extractor_does_not_fake_a_transcript(self):
        """With no CLI available the image is unsupported, never invented."""
        from rfq_copilot.document_extractor import ImageExtractor
        from rfq_copilot.config import Settings

        def missing(argv, **kw):
            raise FileNotFoundError(argv[0])
        s = Settings()
        c = ImageExtractor(s, runner=missing).extract(fixture("supplier_e_quote.png"))
        self.assertIn(c.status, (ExtractionStatus.UNSUPPORTED, ExtractionStatus.FAILED))
        self.assertEqual(c.text, "")

    def test_image_transcripts_are_capped_below_full_confidence(self):
        from rfq_copilot.document_extractor import IMAGE_CONFIDENCE_CEILING
        self.assertLess(IMAGE_CONFIDENCE_CEILING, 1.0, "optical reading is never certain")


class EvidenceGuardTest(unittest.TestCase):
    DOC = ("Item   Size (inch)        Qty       Price per 100 pcs (USD)\n"
           "1      10 x 10 x 5        2000 pcs  41.00\n"
           "2. 5% discount applicable for total order quantities above 10,000 pcs.")

    def test_a_real_span_is_verified(self):
        for span in ("41.00", "10 x 10 x 5", "5% discount applicable"):
            self.assertTrue(sg.evidence_is_findable(span, [self.DOC]), span)

    def test_a_fabricated_span_is_rejected(self):
        for span in ("USD 0.99 per piece", "we offer free shipping", "", "  "):
            self.assertFalse(sg.evidence_is_findable(span, [self.DOC]), span)

    def test_punctuation_and_spacing_differences_still_match(self):
        self.assertTrue(sg.evidence_is_findable("10x10x5", [self.DOC]))

    def test_evidence_records_where_it_came_from(self):
        docs = make_documents(self.DOC, "supplier_b_quote.pdf", "pdf")
        store = {}
        eid = sg.build_evidence({"quoted_text": "41.00", "location": "Page 1"}, "resp_1", docs, store)
        ev = store[eid]
        self.assertTrue(ev.verified)
        self.assertEqual(ev.page, 1)
        self.assertEqual(ev.document_name, "supplier_b_quote.pdf")

    def test_unverifiable_evidence_is_stored_but_flagged(self):
        docs = make_documents(self.DOC)
        store = {}
        eid = sg.build_evidence({"quoted_text": "USD 9.99 per piece", "location": "Page 1"}, "r", docs, store)
        self.assertFalse(store[eid].verified, "a span we cannot find must not pass as evidence")

    def test_a_price_without_findable_evidence_is_held_for_review(self):
        docs = make_documents(self.DOC)
        store = {}
        eid = sg.build_evidence({"quoted_text": "totally made up", "location": "Page 1"}, "r", docs, store)
        q = SupplierQuote(unit_price=0.42, currency="USD", price_basis=PriceBasis.PER_UNIT,
                          match_status=MatchStatus.MATCHED, line_item_id="LINE-001", evidence_ids=[eid])
        sg.guard_quote(q, store)
        self.assertEqual(q.status, QuoteStatus.NEEDS_REVIEW)
        self.assertTrue(any("traced" in i for i in q.issues))


class InventedPriceGuardTest(unittest.TestCase):
    def test_a_price_absent_from_the_document_is_flagged(self):
        docs = make_documents("Line 1  10 x 10 x 5  2000 pcs  0.42")
        real = SupplierQuote(supplier_line_label="L1", unit_price=0.42, currency="USD")
        invented = SupplierQuote(supplier_line_label="L2", unit_price=9.99, currency="USD")
        notes = sg.guard_no_invented_lines([real, invented], docs)
        self.assertEqual(real.status, QuoteStatus.QUOTED)
        self.assertEqual(invented.status, QuoteStatus.NEEDS_REVIEW)
        self.assertEqual(len(notes), 1)

    def test_quotes_for_unknown_rfq_lines_are_discarded(self):
        rfq = carton_rfq()
        q = SupplierQuote(supplier_line_label="mystery", line_item_id="LINE-999",
                          match_status=MatchStatus.MATCHED, unit_price=1.0, currency="USD")
        notes = sg.guard_line_ids([q], rfq)
        self.assertIsNone(q.line_item_id)
        self.assertEqual(q.match_status, MatchStatus.UNMATCHED)
        self.assertTrue(notes)


class CertificationGuardTest(unittest.TestCase):
    def test_a_bare_claim_is_never_verified(self):
        docs = make_documents("We are ISO 9001 certified and operate a documented quality system.")
        store = {}
        eid = sg.build_evidence({"quoted_text": "We are ISO 9001 certified", "location": "Page 3"}, "r", docs, store)
        cert = Certification(raw_name="ISO 9001", evidence_ids=[eid])
        sg.guard_certification(cert, docs, store)
        self.assertEqual(cert.status, ClaimStatus.CLAIMED)
        self.assertEqual(cert.name, "ISO 9001")

    def test_claiming_an_attachment_we_do_not_hold_stays_claimed(self):
        docs = make_documents("ISO 9001 certificate attached.", "supplier_a_quote.xlsx", "xlsx")
        store = {}
        eid = sg.build_evidence({"quoted_text": "ISO 9001 certificate attached", "location": "Sheet Q"}, "r", docs, store)
        cert = Certification(raw_name="ISO 9001", document_reference="certificate attached", evidence_ids=[eid])
        sg.guard_certification(cert, docs, store)
        self.assertEqual(cert.status, ClaimStatus.CLAIMED)
        self.assertIn("not among the documents", cert.note)

    def test_a_held_certificate_document_verifies_the_claim(self):
        docs = make_documents("ISO 9001 certificate.", "iso9001_certificate.pdf", "pdf")
        store = {}
        eid = sg.build_evidence({"quoted_text": "ISO 9001 certificate", "location": "Page 1"}, "r", docs, store)
        cert = Certification(raw_name="ISO 9001", document_reference="iso9001_certificate.pdf", evidence_ids=[eid])
        sg.guard_certification(cert, docs, store)
        self.assertEqual(cert.status, ClaimStatus.VERIFIED)

    def test_a_quotation_cannot_verify_its_own_certificate_claim(self):
        """Found on the live demo: the model reported the quote image as the attachment,
        so the quotation verified the claim it was making. A quote is the claim; it
        cannot also be the proof."""
        docs = make_documents("ISO 9001 certificate attached.", "stress_e_quote.png", "image")
        store = {}
        eid = sg.build_evidence({"quoted_text": "ISO 9001 certificate attached",
                                 "location": "image line 4"}, "r", docs, store)
        cert = Certification(raw_name="ISO 9001", document_reference="stress_e_quote.png",
                             evidence_ids=[eid])
        sg.guard_certification(cert, docs, store, {docs[0].id})
        self.assertEqual(cert.status, ClaimStatus.CLAIMED)
        self.assertIn("their quotation rather than a certificate", cert.note)
        self.assertEqual(cert.document_id, "")

    def test_a_separate_certificate_document_still_verifies(self):
        quote, cert_doc = make_documents("Prices as attached.", "quote.png", "image")[0], \
            make_documents("ISO 9001:2015 — certificate no. TR-9001-4471",
                           "iso_certificate.png", "image")[0]
        store = {}
        eid = sg.build_evidence({"quoted_text": "ISO 9001", "location": "image line 1"},
                                "r", [cert_doc], store)
        cert = Certification(raw_name="ISO 9001", document_reference="iso_certificate.png",
                             evidence_ids=[eid])
        sg.guard_certification(cert, [quote, cert_doc], store, {quote.id})
        self.assertEqual(cert.status, ClaimStatus.VERIFIED)
        self.assertEqual(cert.document_id, cert_doc.id)

    def test_a_document_that_never_mentions_the_certificate_evidences_nothing(self):
        docs = make_documents("Company profile and factory photographs.",
                              "certificates.pdf", "pdf")
        store = {}
        cert = Certification(raw_name="ISO 9001", document_reference="certificates.pdf")
        sg.guard_certification(cert, docs, store)
        self.assertEqual(cert.status, ClaimStatus.CLAIMED)
        self.assertIn("does not mention this certificate", cert.note)

    def test_a_day_first_expiry_is_read_rather_than_ignored(self):
        """"31/12/2020" was unparseable, which left an expired certificate reading as
        current — the one direction that matters."""
        docs = make_documents("ISO 9001 valid until 31/12/2020", "iso_cert.pdf", "pdf")
        store = {}
        eid = sg.build_evidence({"quoted_text": "ISO 9001 valid until 31/12/2020",
                                 "location": "Page 1"}, "r", docs, store)
        cert = Certification(raw_name="ISO 9001", document_reference="iso_cert.pdf",
                             expiry_date="31/12/2020", evidence_ids=[eid])
        sg.guard_certification(cert, docs, store)
        self.assertEqual(cert.status, ClaimStatus.EXPIRED)

    def test_an_expired_certificate_is_marked_expired(self):
        docs = make_documents("ISO 9001 valid until 2020-01-01", "iso.pdf", "pdf")
        store = {}
        eid = sg.build_evidence({"quoted_text": "ISO 9001 valid until 2020-01-01", "location": "Page 1"}, "r", docs, store)
        cert = Certification(raw_name="ISO 9001", document_reference="iso.pdf",
                             expiry_date="2020-01-01", evidence_ids=[eid])
        sg.guard_certification(cert, docs, store)
        self.assertEqual(cert.status, ClaimStatus.EXPIRED)


class ConfidenceTest(unittest.TestCase):
    """A confidence the model actually stated must survive being capped."""

    def test_a_stated_zero_is_not_read_as_missing(self):
        quote = SupplierQuote(unit_price=1.0, currency="USD", confidence=0.0,
                              status=QuoteStatus.QUOTED)
        sg.guard_quote(quote, {}, 1.0)
        self.assertEqual(quote.confidence, 0.0,
                         "0.0 means the model had no confidence, not that it gave none")

    def test_a_missing_confidence_takes_the_ceiling(self):
        quote = SupplierQuote(status=QuoteStatus.NOT_QUOTED, confidence=None)
        sg.guard_quote(quote, {}, 0.75)
        self.assertEqual(quote.confidence, 0.75)

    def test_an_unevidenced_price_is_capped_even_when_the_model_was_sure(self):
        quote = SupplierQuote(unit_price=1.0, currency="USD", confidence=0.99,
                              status=QuoteStatus.QUOTED)
        sg.guard_quote(quote, {}, 1.0)
        self.assertEqual(quote.confidence, 0.4)


class QuestionnaireGuardTest(unittest.TestCase):
    def test_answers_must_map_to_a_real_question(self):
        rfq = carton_rfq()
        docs = make_documents("Payment: 30% advance.")
        store = {}
        good = QuestionnaireResponse(field_key="payment_terms", answer="30% advance")
        bogus = QuestionnaireResponse(field_key="not_a_field", answer="something")
        kept, notes = sg.guard_questionnaire([good, bogus], rfq, docs, store)
        self.assertEqual([a.field_key for a in kept], ["payment_terms"])
        self.assertEqual(kept[0].question_id, "q_pay", "the stable question id must be attached")
        self.assertTrue(notes)

    def test_an_unanswered_item_is_missing_not_no(self):
        rfq = carton_rfq()
        answered = [QuestionnaireResponse(field_key="payment_terms", answer="30% advance")]
        missing = sg.missing_questionnaire_items(answered, rfq)
        keys = {m.field_key for m in missing}
        self.assertIn("certifications", keys)
        self.assertTrue(all(m.status == ClaimStatus.MISSING for m in missing))
        self.assertTrue(all(m.answer == "" for m in missing), "silence is never turned into an answer")

    def test_a_yes_is_a_claim_not_a_verification(self):
        rfq = carton_rfq()
        docs = make_documents("Yes, we are ISO 9001 certified.")
        store = {}
        eid = sg.build_evidence({"quoted_text": "Yes, we are ISO 9001 certified", "location": "line 1"}, "r", docs, store)
        a = QuestionnaireResponse(field_key="certifications", answer="Yes, we are ISO 9001 certified",
                                  evidence_ids=[eid])
        kept, _ = sg.guard_questionnaire([a], rfq, docs, store)
        self.assertEqual(kept[0].status, ClaimStatus.CLAIMED)


class ConflictAndRevisionTest(unittest.TestCase):
    def test_both_sides_of_a_conflict_are_kept(self):
        docs = make_documents("lead time of 15 days\nlead time of 25 days should be allowed")
        store = {}
        raw = [{"topic": "lead time", "description": "Two different lead times",
                "values": [{"value": "15 days", "evidence": {"quoted_text": "lead time of 15 days", "location": "Page 1"}},
                           {"value": "25 days", "evidence": {"quoted_text": "lead time of 25 days", "location": "Page 3"}}]}]
        out = sg.record_conflicts(raw, "r", docs, store)
        self.assertEqual(len(out), 1)
        self.assertEqual(len(out[0]["values"]), 2, "a conflict is never collapsed to one value")
        self.assertTrue(all(v["evidence_id"] for v in out[0]["values"]))

    def test_a_one_sided_conflict_is_not_a_conflict(self):
        docs = make_documents("lead time 15 days")
        out = sg.record_conflicts([{"topic": "x", "description": "", "values": [{"value": "15", "evidence": None}]}],
                                  "r", docs, {})
        self.assertEqual(out, [])

    def test_the_latest_response_becomes_active_and_earlier_ones_survive(self):
        from rfq_copilot.supplier_models import SupplierResponse
        first = SupplierResponse(id="resp_1", received_at="2026-09-09")
        second = SupplierResponse(id="resp_2", received_at="2026-09-12")
        sg.resolve_revision_chain([first, second])
        self.assertFalse(first.is_active)
        self.assertTrue(second.is_active)
        self.assertEqual(first.superseded_by_id, "resp_2")
        self.assertEqual(second.revises_response_id, "resp_1")


class CurrencyGuardTest(unittest.TestCase):
    def _quote(self, currency):
        q = SupplierQuote(unit_price=41.0, currency=currency, price_basis=PriceBasis.PER_UNIT,
                          normalized_unit_price=41.0, normalization_status=NormalizationStatus.NORMALIZED)
        return sg.guard_currency(q)

    def test_known_codes_and_unambiguous_symbols_pass(self):
        for given, expected in (("USD", "USD"), ("usd", "USD"), ("EUR", "EUR"), ("€", "EUR"), ("₹", "INR")):
            q = self._quote(given)
            self.assertEqual(q.currency, expected)
            self.assertEqual(q.status, QuoteStatus.QUOTED)
            self.assertIsNotNone(q.normalized_unit_price)

    def test_a_subunit_without_its_currency_is_not_a_price(self):
        """'41 cents' could be a US cent or an Indian paisa; we refuse to choose."""
        q = self._quote("cents")
        self.assertEqual(q.status, QuoteStatus.NEEDS_REVIEW)
        self.assertIsNone(q.normalized_unit_price, "an unnamed currency must not reach the price column")
        self.assertEqual(q.normalization_status, NormalizationStatus.UNRESOLVED)

    def test_a_bare_dollar_sign_is_ambiguous(self):
        q = self._quote("$")
        self.assertEqual(q.status, QuoteStatus.NEEDS_REVIEW)
        self.assertIsNone(q.normalized_unit_price)

    def test_a_missing_currency_blocks_comparison(self):
        q = self._quote("")
        self.assertEqual(q.status, QuoteStatus.NEEDS_REVIEW)
        self.assertIsNone(q.normalized_unit_price)


if __name__ == "__main__":
    unittest.main()
