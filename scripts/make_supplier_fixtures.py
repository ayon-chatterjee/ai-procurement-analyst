"""Generate the five messy supplier fixtures used by the Phase 2 demo.

These are *source documents*, not extraction results: real .xlsx, .pdf, .docx, .txt and
.png files that the DocumentExtractor has to read for itself. Each one is deliberately
awkward in a different, realistic way, so the extraction pipeline is exercised rather
than flattered:

  A  clean spreadsheet, complete quote, everything in tidy columns
  B  PDF priced per 100 pieces, discount hidden in a footnote, lead time stated twice
     with different numbers, MOQ above the RFQ quantity  (+ a later revision)
  C  Word prose, partial coverage, one line priced per kilogram, validity conditional,
     and a question back to the buyer
  D  plain-text email referring to sizes by position, prices in cents, lines missing
  E  photographed quotation in EUR, angled and noisy, to be read by vision

Run:  python3 scripts/make_supplier_fixtures.py
"""
from __future__ import annotations

import os
import sys
import zipfile

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(ROOT, "fixtures", "suppliers")

# The seven carton sizes of the hero RFQ (inches), 2,000 pieces each.
SIZES = ["10 x 10 x 5", "12 x 10 x 6", "15 x 10 x 8", "18 x 12 x 10",
         "20 x 15 x 10", "24 x 18 x 12", "30 x 20 x 15"]


# --------------------------------------------------------------------------- #
# A — clean Excel
# --------------------------------------------------------------------------- #
def supplier_a() -> str:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active
    ws.title = "Quotation"
    ws["A1"] = "ANHUI PACKAGING CO., LTD"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = "Quotation Ref: AP-2026-0418     Date: 2026-09-08"
    ws["A3"] = "Attn: Procurement — Corrugated Carton Boxes RFQ"

    headers = ["Line", "Description", "Qty", "Unit", "Unit Price (USD)", "Currency"]
    for col, h in enumerate(headers, start=1):
        c = ws.cell(row=5, column=col, value=h)
        c.font = Font(bold=True)

    prices = [0.42, 0.48, 0.55, 0.71, 0.84, 1.05, 1.38]
    for i, (size, price) in enumerate(zip(SIZES, prices)):
        r = 6 + i
        ws.cell(row=r, column=1, value=i + 1)
        ws.cell(row=r, column=2, value="%s inch corrugated carton" % size)
        ws.cell(row=r, column=3, value=2000)
        ws.cell(row=r, column=4, value="pcs")
        ws.cell(row=r, column=5, value=price)
        ws.cell(row=r, column=6, value="USD")

    ws["A15"] = "Commercial Terms"
    ws["A15"].font = Font(bold=True)
    ws["A16"] = "Minimum order quantity"
    ws["B16"] = "2,000 pcs per size"
    ws["A17"] = "Production lead time"
    ws["B17"] = "18 days from artwork approval"
    ws["A18"] = "Payment terms"
    ws["B18"] = "30% advance, 70% against B/L copy"
    ws["A19"] = "Price basis"
    ws["B19"] = "FOB Shanghai"
    ws["A20"] = "Quotation validity"
    ws["B20"] = "30 days from date of issue"

    ws["A22"] = "Certifications"
    ws["A22"].font = Font(bold=True)
    ws["A23"] = "ISO 9001:2015"
    ws["B23"] = "Cert No. CN-ISO-884512, issued by SGS, valid until 2027-03-14"
    ws["A24"] = "FSC Chain of Custody"
    ws["B24"] = "Cert No. FSC-C129945, valid until 2026-11-30"
    ws["A25"] = "Certificates attached to this quotation."

    for col, width in (("A", 26), ("B", 44), ("C", 10), ("D", 8), ("E", 18), ("F", 10)):
        ws.column_dimensions[col].width = width

    path = os.path.join(OUT, "supplier_a_quote.xlsx")
    wb.save(path)
    return path


# --------------------------------------------------------------------------- #
# PDF writing (no dependency: a small, valid, uncompressed PDF)
# --------------------------------------------------------------------------- #
def _pdf_escape(text: str) -> str:
    return text.replace("\\", r"\\").replace("(", r"\(").replace(")", r"\)")


def write_pdf(path: str, pages: list) -> str:
    """pages: list of lists of text lines. Produces a real PDF with a correct xref."""
    objects = {}
    font_id = 100
    objects[font_id] = b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"

    page_ids, content_ids = [], []
    for n, lines in enumerate(pages):
        pid, cid = 10 + n * 2, 11 + n * 2
        page_ids.append(pid)
        content_ids.append(cid)
        body = ["BT", "/F1 11 Tf", "14 TL", "50 760 Td"]
        for line in lines:
            body.append("(%s) Tj T*" % _pdf_escape(line))
        body.append("ET")
        stream = "\n".join(body).encode("latin-1", "replace")
        objects[cid] = b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream"

    kids = " ".join("%d 0 R" % p for p in page_ids)
    objects[1] = b"<< /Type /Catalog /Pages 2 0 R >>"
    objects[2] = ("<< /Type /Pages /Kids [%s] /Count %d >>" % (kids, len(page_ids))).encode()
    for pid, cid in zip(page_ids, content_ids):
        objects[pid] = ("<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
                        "/Contents %d 0 R /Resources << /Font << /F1 %d 0 R >> >> >>" % (cid, font_id)).encode()

    out = bytearray(b"%PDF-1.4\n")
    offsets = {}
    for num in sorted(objects):
        offsets[num] = len(out)
        out += b"%d 0 obj\n" % num + objects[num] + b"\nendobj\n"

    max_num = max(objects) + 1
    xref_at = len(out)
    out += b"xref\n0 %d\n" % max_num
    out += b"0000000000 65535 f \n"
    for num in range(1, max_num):
        if num in offsets:
            out += b"%010d 00000 n \n" % offsets[num]
        else:
            out += b"0000000000 65535 f \n"
    out += b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (max_num, xref_at)

    with open(path, "wb") as f:
        f.write(bytes(out))
    return path


def supplier_b() -> str:
    prices = [41.0, 47.0, 54.0, 69.0, 82.0, 103.0, 136.0]     # per 100 pieces
    page1 = [
        "SHENZHEN PRINT & PACK LIMITED",
        "Quotation SZPP/Q/2026/1147          Date: 09 September 2026",
        "",
        "RE: Corrugated carton boxes - request for quotation",
        "",
        "All prices below are quoted PER 100 PIECES, EXW Shenzhen.",
        "",
        "Item   Size (inch)        Qty       Price per 100 pcs (USD)",
    ]
    for i, (size, p) in enumerate(zip(SIZES, prices)):
        page1.append("%-6d %-18s %-9s %.2f" % (i + 1, size, "2000 pcs", p))
    page1 += [
        "",
        "Minimum order quantity: 5,000 pieces per size.",
        "Production lead time: 15 days after receipt of approved artwork.",
        "Payment: T/T 50% deposit, balance before shipment.",
    ]
    page2 = [
        "SHENZHEN PRINT & PACK LIMITED - page 2",
        "",
        "Board specification: 3-ply B-flute, 125gsm kraft liner.",
        "Printing: up to 2 colours flexo included in the above prices.",
        "",
        "Quotation is valid for 21 days from the date of issue.",
        "",
        "Notes:",
        "1. Prices exclude export carton pallets.",
        "2. 5% discount applicable for total order quantities above 10,000 pcs.",
        "3. Tooling/die charges waived for repeat orders.",
    ]
    page3 = [
        "SHENZHEN PRINT & PACK LIMITED - page 3",
        "",
        "Quality and compliance",
        "",
        "We are ISO 9001 certified and operate a documented quality system.",
        "BRC packaging certification is in progress.",
        "",
        "Delivery schedule: please note production lead time of 25 days",
        "should be allowed during peak season (September to November).",
        "",
        "For and on behalf of Shenzhen Print & Pack Limited",
    ]
    return write_pdf(os.path.join(OUT, "supplier_b_quote.pdf"), [page1, page2, page3])


def supplier_b_revision() -> str:
    prices = [39.0, 45.0, 52.0, 66.0, 79.0, 99.0, 131.0]
    page1 = [
        "SHENZHEN PRINT & PACK LIMITED",
        "REVISED QUOTATION SZPP/Q/2026/1147-R1      Date: 12 September 2026",
        "",
        "This revision supersedes our quotation dated 09 September 2026.",
        "",
        "Revised prices, still quoted PER 100 PIECES, EXW Shenzhen:",
        "",
        "Item   Size (inch)        Qty       Price per 100 pcs (USD)",
    ]
    for i, (size, p) in enumerate(zip(SIZES, prices)):
        page1.append("%-6d %-18s %-9s %.2f" % (i + 1, size, "2000 pcs", p))
    page1 += [
        "",
        "Minimum order quantity reduced to 3,000 pieces per size.",
        "Production lead time: 15 days after receipt of approved artwork.",
        "The 5% discount above 10,000 pcs total continues to apply.",
        "Quotation is valid for 21 days from the date of issue.",
    ]
    return write_pdf(os.path.join(OUT, "supplier_b_revision.pdf"), [page1])


# --------------------------------------------------------------------------- #
# C — Word prose
# --------------------------------------------------------------------------- #
DOCX_CONTENT_TYPES = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">
<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>
<Default Extension="xml" ContentType="application/xml"/>
<Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>
</Types>"""

DOCX_RELS = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/>
</Relationships>"""


def write_docx(path: str, paragraphs: list) -> str:
    def esc(t):
        return t.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
    body = "".join(
        '<w:p><w:r><w:t xml:space="preserve">%s</w:t></w:r></w:p>' % esc(p) for p in paragraphs)
    document = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
                '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
                '<w:body>%s</w:body></w:document>' % body)
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("[Content_Types].xml", DOCX_CONTENT_TYPES)
        z.writestr("_rels/.rels", DOCX_RELS)
        z.writestr("word/document.xml", document)
    return path


def supplier_c() -> str:
    paras = [
        "VIET CARTON JOINT STOCK COMPANY",
        "Response to your enquiry for corrugated cartons - 10 September 2026",
        "",
        "Dear Procurement team,",
        "Thank you for the enquiry. We are pleased to offer as follows, although we are "
        "not able to cover every size in your list at this time.",
        "",
        "For the 12 x 10 x 6 box, our price is USD 0.47 per piece for orders above 2,000 units. "
        "Production lead time is approximately 18 days from confirmation of artwork.",
        "",
        "The smaller 10 x 10 x 5 carton we can supply at USD 0.44 per piece on the same terms.",
        "",
        "For the large carton, 24 x 18 x 12, we would quote USD 1.02 per piece.",
        "",
        "We can also offer the 30 x 20 x 15 size, but for this item our mill prices by weight: "
        "the price would be USD 2.35 per kg of finished carton. We can confirm a per piece "
        "figure once the final board grade is agreed.",
        "",
        "Regrettably we cannot quote the 15 x 10 x 8, 18 x 12 x 10 and 20 x 15 x 10 sizes in "
        "this round as our die inventory does not cover them.",
        "",
        "Our minimum order quantity is 2,000 pieces per size. Payment terms are 40% deposit "
        "with the balance against shipping documents. Delivery would be FOB Haiphong.",
        "",
        "Please note that this quotation is valid subject to kraft paper prices, which have "
        "been volatile this quarter. We would confirm final pricing at the time of order.",
        "",
        "One clarification please: could you confirm whether 5-colour printing is required for "
        "all sizes, or only for the retail-facing cartons? This materially affects our tooling cost.",
        "",
        "We hold ISO 9001 certification. A copy of the certificate can be provided on request.",
        "",
        "Best regards,",
        "Nguyen Thi Mai, Export Sales, Viet Carton JSC",
    ]
    return write_docx(os.path.join(OUT, "supplier_c_response.docx"), paras)


# --------------------------------------------------------------------------- #
# D — plain text email
# --------------------------------------------------------------------------- #
def supplier_d() -> str:
    text = """From: sales@gujaratboxes.in
To: procurement@buyer.example
Subject: Re: RFQ corrugated cartons
Date: Thu, 10 Sep 2026 16:12:04 +0530

Hi,

Thanks for your enquiry. We can do sizes 1, 2, 3 and 5 from your list.

Price is 41 cents, 49 cents, 62 cents and 88 cents respectively, per piece,
ex works Ahmedabad.

MOQ 1000 each. Lead time 21 days. ISO available.

We are not able to offer sizes 4, 6 and 7 at present.

Regards,
Rakesh Patel
Gujarat Boxes Pvt Ltd
"""
    path = os.path.join(OUT, "supplier_d_response.txt")
    with open(path, "w", encoding="utf-8") as f:
        f.write(text)
    return path


# --------------------------------------------------------------------------- #
# E — photographed quotation (read by vision, not by a text parser)
# --------------------------------------------------------------------------- #
def supplier_e() -> str:
    from PIL import Image, ImageDraw, ImageFilter

    W, H = 1200, 1000
    canvas = Image.new("RGB", (W, H), (232, 230, 226))
    page = Image.new("RGB", (980, 820), (253, 252, 249))
    d = ImageDraw.Draw(page)

    def line(x, y, text, size=1):
        # default PIL font is small; draw twice with a 1px offset for a bolder look
        d.text((x, y), text, fill=(22, 22, 26))
        if size > 1:
            d.text((x + 1, y), text, fill=(22, 22, 26))

    line(40, 34, "ISTANBUL AMBALAJ SANAYI A.S.", 2)
    line(40, 58, "Teklif / Quotation  No: IA-2026-0912")
    line(40, 76, "Tarih / Date: 11.09.2026")
    d.line([(40, 98), (940, 98)], fill=(90, 90, 95), width=2)

    line(40, 118, "QUOTATION - CORRUGATED CARTONS      All prices in EUR per piece")
    line(40, 150, "No.   Size (inch)          Qty        Unit Price (EUR)")
    d.line([(40, 168), (700, 168)], fill=(150, 150, 155), width=1)

    rows = [("1", "10 x 10 x 5", "2000", "0.39"),
            ("2", "12 x 10 x 6", "2000", "0.44"),
            ("4", "18 x 12 x 10", "2000", "0.66"),
            ("5", "20 x 15 x 10", "2000", "0.78"),
            ("6", "24 x 18 x 12", "2000", "0.97")]
    y = 186
    for no, size, qty, price in rows:
        line(46, y, no)
        line(110, y, size)
        line(300, y, qty)
        line(430, y, price)
        y += 30

    y += 18
    line(40, y, "Sizes 3 and 7 not quoted - tooling not available.")
    line(40, y + 30, "MOQ: 2000 pcs per size")
    line(40, y + 54, "Lead time: 24 days after order confirmation")
    line(40, y + 78, "Payment: 50% advance, 50% before dispatch")
    line(40, y + 102, "Delivery: FOB Istanbul")
    line(40, y + 126, "Validity: 15 days")
    line(40, y + 162, "ISO 9001 certificate no. TR-9001-4471 attached.")
    line(40, y + 200, "Kind regards,")
    line(40, y + 224, "Emre Yilmaz - Sales Manager")

    # make it look photographed: slight rotation, soft focus, uneven lighting
    page = page.rotate(-2.4, expand=True, fillcolor=(232, 230, 226))
    canvas.paste(page, (100, 70))
    canvas = canvas.filter(ImageFilter.GaussianBlur(0.6))
    shade = Image.linear_gradient("L").resize((W, H)).point(lambda v: 218 + v // 8)
    canvas = Image.composite(canvas, Image.new("RGB", (W, H), (255, 255, 255)), shade)

    path = os.path.join(OUT, "supplier_e_quote.png")
    canvas.save(path, "PNG")
    return path


# --------------------------------------------------------------------------- #
def main() -> int:
    os.makedirs(OUT, exist_ok=True)
    made = [supplier_a(), supplier_b(), supplier_b_revision(), supplier_c(), supplier_d(), supplier_e()]
    for p in made:
        print("  %-34s %8d bytes" % (os.path.basename(p), os.path.getsize(p)))
    print("\nfixtures written to %s" % OUT)
    return 0


if __name__ == "__main__":
    sys.exit(main())
